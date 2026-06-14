"""
CSV / XLSX → ChromaDB インジェスト
CSV行番号ベースIDによる差分更新対応
"""
import argparse
import csv
import curses
from pathlib import Path
from typing import Callable

import openpyxl

import chromadb
from langchain_chroma import Chroma
from langchain_core.documents import Document
from langchain_ollama import OllamaEmbeddings

BASE_DIR = Path.cwd()
CSV_GLOB = "引き継ぎノート*.csv"
XLSX_GLOB = "引き継ぎノート*.xlsx"
CHROMA_PATH = str(Path(__file__).parent / "chroma_db")
COLLECTION_NAME = "hikitsugi"
EMBED_MODEL = "nomic-embed-text"
BATCH_SIZE = 50


def _curses_select(files: list[Path]) -> Path | None:
    """カーソルキーで .csv/.xlsx ファイルを選択する。Escまたは'q'でキャンセル。"""
    def _inner(stdscr: "curses._CursesWindow") -> Path | None:
        curses.curs_set(0)
        curses.init_pair(1, curses.COLOR_BLACK, curses.COLOR_WHITE)
        idx = 0
        while True:
            stdscr.clear()
            h, w = stdscr.getmaxyx()
            stdscr.addstr(0, 0, "ファイルを選択してください (↑↓: 移動, Enter: 決定, q: キャンセル)")
            for i, f in enumerate(files):
                y = i + 2
                if y >= h - 1:
                    break
                label = f.name[:w - 4]
                if i == idx:
                    stdscr.attron(curses.color_pair(1))
                    stdscr.addstr(y, 2, label)
                    stdscr.attroff(curses.color_pair(1))
                else:
                    stdscr.addstr(y, 2, label)
            stdscr.refresh()
            key = stdscr.getch()
            if key in (curses.KEY_UP, ord("k")) and idx > 0:
                idx -= 1
            elif key in (curses.KEY_DOWN, ord("j")) and idx < len(files) - 1:
                idx += 1
            elif key in (curses.KEY_ENTER, ord("\n"), ord("\r")):
                return files[idx]
            elif key in (ord("q"), 27):  # q or Esc
                return None

    return curses.wrapper(_inner)


def select_file_interactively() -> Path | None:
    """カレントディレクトリの .csv/.xlsx をカーソル選択させる。ファイルがなければ None。"""
    files = sorted(
        list(BASE_DIR.glob("*.csv")) + list(BASE_DIR.glob("*.xlsx")),
        key=lambda p: p.name,
    )
    if not files:
        return None
    return _curses_select(files)


def find_default_file() -> Path | None:
    """引き継ぎノート*.csv / *.xlsx を BASE_DIR から検索し、最終更新日時が最新のものを返す。"""
    candidates = sorted(
        list(BASE_DIR.glob(CSV_GLOB)) + list(BASE_DIR.glob(XLSX_GLOB)),
        key=lambda p: p.stat().st_mtime,
        reverse=True,
    )
    return candidates[0] if candidates else None


# 後方互換エイリアス
find_default_csv = find_default_file


def _parse_row(row: list[str], i: int) -> tuple[str, Document] | None:
    """行データから (doc_id, Document) を生成。CSVとXLSX共通ロジック。"""
    if len(row) < 8:
        return None
    # row[0]: 空欄, row[1]: 行番号（無視）
    content = row[7].strip().replace("　", " ")
    if not content:
        return None

    date = row[2].strip()
    youbi = row[3].strip()
    kinmu = row[4].strip()
    shubetsu = row[5].strip()
    jikoku = row[6].strip()
    author = row[8].strip() if len(row) > 8 else ""

    parts = []
    if date:
        parts.append(f"日付: {date}({youbi})")
    if kinmu:
        parts.append(f"勤務: {kinmu}")
    if shubetsu:
        parts.append(f"種別: {shubetsu}")
    if jikoku:
        parts.append(f"時刻: {jikoku}")
    parts.append(f"内容: {content}")
    if author:
        parts.append(f"記入者: {author}")

    return (f"csv_row_{i}", Document(
        page_content="\n".join(parts),
        metadata={
            "日付": date,
            "曜日": youbi,
            "勤務": kinmu,
            "種別": shubetsu,
            "時刻": jikoku,
            "記入者": author,
            "row_index": i,
        },
    ))


def load_documents_from_csv(csv_path: Path) -> list[tuple[str, Document]]:
    """CSVを読んで (doc_id, Document) のリストを返す。IDはCSVデータ行番号ベース。"""
    results = []
    with open(csv_path, encoding="utf-8-sig") as f:
        reader = csv.reader(f)
        next(reader)  # タイトル行
        next(reader)  # ヘッダー行
        for i, row in enumerate(reader):
            pair = _parse_row([c.strip() for c in row], i)
            if pair:
                results.append(pair)
    return results


def load_documents_from_xlsx(xlsx_path: Path) -> list[tuple[str, Document]]:
    """XLSXを読んで (doc_id, Document) のリストを返す。IDはデータ行番号ベース。"""
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb.active
    results = []
    rows = ws.iter_rows(values_only=True)
    next(rows)  # タイトル行
    next(rows)  # ヘッダー行
    for i, raw_row in enumerate(rows):
        row = [str(cell) if cell is not None else "" for cell in raw_row]
        pair = _parse_row(row, i)
        if pair:
            results.append(pair)
    wb.close()
    return results


def load_documents(path: Path) -> list[tuple[str, Document]]:
    """拡張子に応じてCSVまたはXLSXからドキュメントを読み込む。"""
    if path.suffix.lower() == ".xlsx":
        return load_documents_from_xlsx(path)
    return load_documents_from_csv(path)


def run_ingest(
    csv_path: Path,
    progress_cb: Callable[[int, int, str], None] | None = None,
    rebuild: bool = False,
) -> dict:
    """
    CSVをChromaDBに差分インジェスト。
    progress_cb(done_count, total_count, message) で進捗を通知。
    rebuild=True の場合はコレクションを削除して全件再構築。
    戻り値: {"added": int, "skipped": int, "total": int}
    """
    def cb(done: int, total: int, msg: str) -> None:
        if progress_cb:
            progress_cb(done, total, msg)

    cb(0, 0, f"ファイル読み込み中: {csv_path.name}")
    id_doc_pairs = load_documents(csv_path)
    cb(0, 0, f"有効レコード: {len(id_doc_pairs)} 件")

    client = chromadb.PersistentClient(path=CHROMA_PATH)
    existing_names = [c.name for c in client.list_collections()]

    if rebuild and COLLECTION_NAME in existing_names:
        cb(0, 0, "既存コレクションを削除中...")
        client.delete_collection(COLLECTION_NAME)
        existing_ids: set[str] = set()
    elif COLLECTION_NAME in existing_names:
        col = client.get_collection(COLLECTION_NAME)
        all_existing = col.get(include=[])
        existing_ids = set(all_existing["ids"])
        cb(0, 0, f"既存DB: {len(existing_ids)} 件登録済み")
    else:
        existing_ids = set()

    new_pairs = [(id_, doc) for id_, doc in id_doc_pairs if id_ not in existing_ids]
    skipped = len(id_doc_pairs) - len(new_pairs)

    if not new_pairs:
        total_in_db = len(existing_ids)
        cb(0, 0, f"新規レコードなし（{skipped} 件スキップ）")
        return {"added": 0, "skipped": skipped, "total": total_in_db}

    cb(0, len(new_pairs), f"新規 {len(new_pairs)} 件を埋め込み中...")

    embeddings = OllamaEmbeddings(model=EMBED_MODEL)
    db = Chroma(
        collection_name=COLLECTION_NAME,
        embedding_function=embeddings,
        persist_directory=CHROMA_PATH,
    )

    added = 0
    for i in range(0, len(new_pairs), BATCH_SIZE):
        batch = new_pairs[i:i + BATCH_SIZE]
        db.add_documents(
            documents=[p[1] for p in batch],
            ids=[p[0] for p in batch],
        )
        added += len(batch)
        cb(added, len(new_pairs), f"{added}/{len(new_pairs)} 件完了")

    db = None

    client2 = chromadb.PersistentClient(path=CHROMA_PATH)
    total = client2.get_collection(COLLECTION_NAME).count()

    cb(added, added, f"完了: DB合計 {total} 件")
    return {"added": added, "skipped": skipped, "total": total}


def main() -> None:
    parser = argparse.ArgumentParser(description="CSV/XLSXをChromaDBにインジェスト")
    parser.add_argument(
        "--csv", type=Path, default=None,
        help=f"CSV/XLSXファイルパス (省略時は {BASE_DIR}/{CSV_GLOB} または {XLSX_GLOB} を自動検索)",
    )
    parser.add_argument("--rebuild", action="store_true", help="全件再構築（既存を削除）")
    args = parser.parse_args()

    csv_path = args.csv or find_default_file()
    if csv_path is None:
        print(f"カレントディレクトリに {CSV_GLOB} / {XLSX_GLOB} が見つかりませんでした。")
        csv_path = select_file_interactively()
        if csv_path is None:
            # カレントに候補がないか、選択をキャンセルした場合は手動入力へ
            while True:
                user_input = input("CSV/XLSXファイルのパスを入力してください: ").strip()
                if not user_input:
                    print("パスが入力されていません。終了します。")
                    raise SystemExit(1)
                candidate = Path(user_input)
                if not candidate.exists():
                    print(f"ファイルが存在しません: {candidate}")
                    continue
                if candidate.suffix.lower() not in (".csv", ".xlsx"):
                    print(f"対応していない拡張子です（.csv / .xlsx のみ）: {candidate.suffix}")
                    continue
                csv_path = candidate
                break

    result = run_ingest(csv_path, lambda d, t, m: print(m), rebuild=args.rebuild)
    print(result)


if __name__ == "__main__":
    main()
